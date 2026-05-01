import os
import ast
from datetime import datetime
from docx import Document

ROOT = "."
REPORT_ROOT = "scan_reports"


def scan_python_files(root):
    py_files = []
    for path, _, files in os.walk(root):
        for f in files:
            if f.endswith(".py"):
                py_files.append(os.path.join(path, f))
    return py_files


def parse_imports(file):
    imports = []
    try:
        with open(file, "r", encoding="utf-8") as f:
            tree = ast.parse(f.read())

        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                for n in node.names:
                    imports.append(n.name)

            elif isinstance(node, ast.ImportFrom):
                if node.module:
                    imports.append(node.module)

    except Exception as e:
        imports.append(f"PARSE_ERROR: {str(e)}")

    return imports


def build_dependency_graph():
    graph = {}
    files = scan_python_files(ROOT)

    for f in files:
        name = os.path.basename(f)
        imports = parse_imports(f)
        graph[name] = imports

    return graph


def create_report_folder():
    now = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    folder = os.path.join(REPORT_ROOT, now)
    os.makedirs(folder, exist_ok=True)
    return folder, now


def save_word_report(graph, folder, timestamp):

    doc = Document()
    doc.add_heading("AI Agent Architecture Scan", level=1)
    doc.add_paragraph(f"Scan Time: {timestamp}")

    for file, deps in graph.items():
        doc.add_heading(file, level=2)
        if not deps:
            doc.add_paragraph("No imports detected")
        else:
            for d in deps:
                doc.add_paragraph(d, style="List Bullet")

    file_path = os.path.join(folder, "dependency_report.docx")
    doc.save(file_path)

    print(f"\nReport saved: {file_path}")


if __name__ == "__main__":

    print("Scanning agent architecture...")
    graph = build_dependency_graph()

    folder, ts = create_report_folder()
    save_word_report(graph, folder, ts)

    print("Scan complete.")